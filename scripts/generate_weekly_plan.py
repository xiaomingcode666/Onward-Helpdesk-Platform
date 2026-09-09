from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT = Path(__file__).resolve().parents[1] / "outputs" / "Onward_本周任务计划与完成记录_20260908.xlsx"
SOURCE = r"D:\桌面\Onward_业务功能对照与实施排期_修订讨论稿_20260908.xlsx"

NAVY = "243447"
TEAL = "0F6B78"
GREEN = "DDEFE2"
GREEN_TEXT = "196B3A"
AMBER = "FFF2CC"
RED = "FCE4E4"
LIGHT = "F3F6F8"
WHITE = "FFFFFF"
GRAY = "59636E"
THIN = Side(style="thin", color="D8DEE5")


def base_sheet(ws, title, subtitle, width):
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=18, bold=True, color=WHITE)
    ws["A1"].fill = PatternFill("solid", fgColor=NAVY)
    ws["A1"].alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
    ws["A2"] = subtitle
    ws["A2"].font = Font(name="Arial", size=10, color=GRAY)
    ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 30


def style_header(ws, row, start, end):
    for current in ws[row][start - 1:end]:
        current.font = Font(name="Arial", size=10, bold=True, color=WHITE)
        current.fill = PatternFill("solid", fgColor=TEAL)
        current.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        current.border = Border(bottom=THIN)
    ws.row_dimensions[row].height = 30


def style_body(ws, start_row, end_row, start_col, end_col):
    for row in ws.iter_rows(min_row=start_row, max_row=end_row, min_col=start_col, max_col=end_col):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="20262D")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=THIN)
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill("solid", fgColor="FAFBFC")


def add_table(ws, ref, name):
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def set_widths(ws, widths):
    for column, width in widths.items():
        ws.column_dimensions[column].width = width


def build_workbook():
    wb = Workbook()
    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    ws = wb.active
    ws.title = "本周任务"
    base_sheet(
        ws,
        "Onward 本周任务计划与完成记录",
        "本周口径：2026-09-08 至 2026-09-11（剩余工作日）。范围来自修订讨论稿中仅有的 3 个“待验证”功能项。",
        11,
    )

    summary = [
        ("A3", "任务数", "B3", "=COUNTA(A7:A11)"),
        ("D3", "已完成", "E3", '=COUNTIF(H7:H11,"已完成")'),
        ("G3", "计划人日", "H3", "=SUM(G7:G11)"),
        ("J3", "未完成", "K3", '=COUNTIF(H7:H11,"<>已完成")'),
    ]
    for label_cell, label, value_cell, formula in summary:
        ws[label_cell] = label
        ws[label_cell].font = Font(name="Arial", bold=True, color=GRAY)
        ws[value_cell] = formula
        ws[value_cell].font = Font(name="Arial", size=14, bold=True, color=NAVY)
        ws[value_cell].alignment = Alignment(horizontal="left")
    ws["H3"].number_format = '0.0 "人日"'
    ws["B3"].comment = Comment(
        "任务拆分依据为原工作簿的 3 个“待验证”项，另含排期分析和最终回归。", "Codex"
    )
    ws["H3"].comment = Comment(
        "按 1 名开发者本周剩余 4 个工作日估算；不等同于原表每项 3 人日的跨角色团队估算。", "Codex"
    )

    headers = [
        "任务编号", "计划日期", "功能ID", "业务功能", "本周工作",
        "完成标准", "计划人日", "状态", "完成结果", "验证证据", "业务验收边界",
    ]
    ws.append([])
    ws.append([])
    ws.append(headers)
    tasks = [
        [
            "PLAN-001", date(2026, 9, 8), "", "排期熟悉与范围冻结",
            "解析 7 张工作表、90 个功能项，交叉检查 S0、优先级、外部依赖与仓库现状。",
            "形成可执行的本周范围，明确不具备业务验收条件的事项。",
            0.5, "已完成",
            "确认本周只承诺 FUN-021、FUN-028、FUN-083 的代码复验与缺陷修复。",
            f"来源：{SOURCE}",
            "原表日期和 889 人日尚未团队确认；本周计划不是正式上线承诺。",
        ],
        [
            "PLAN-002", date(2026, 9, 9), "FUN-021", "到期自动关闭",
            "核查策略开关、关闭窗口、状态条件、例外、审计记录和重开链路；补齐回归测试。",
            "未到期、未解决、安全关键、未结束会议及活跃供应商协作不关闭；到期项关闭且历史可追溯、可重开。",
            1.0, "已完成",
            "统一批次时间基准；新增策略关闭、窗口、状态、供应商协作、审计和重开覆盖。",
            "Go 服务层验收：PASS。",
            "自动关闭天数、业务例外及关闭前提醒仍需业务负责人批准。",
        ],
        [
            "PLAN-003", date(2026, 9, 10), "FUN-083", "上传文件类型和大小限制",
            "核查所有上传入口是否统一校验；补异常 ZIP、压缩率、目录穿越和包内危险文件保护。",
            "实际大小和 MIME 不信任客户端；危险内容及异常压缩包被拒绝；正常 ZIP/Office 文件可通过。",
            1.5, "已完成",
            "新增可配置的 ZIP 条目数、解压总量和压缩率限制，并校验加密条目、路径及包内扩展名。",
            "Go 上传安全与企业上传集成回归：PASS。",
            "隔离区和病毒处置证据属于 FUN-084；生产 ClamAV 与批准的类型清单仍待确认。",
        ],
        [
            "PLAN-004", date(2026, 9, 11), "FUN-028", "产品支持组及工程师配置",
            "复验产品组自动建立、Owner、成员变更权限、跨组隔离、分派与多产品组范围。",
            "产品可配置支持组并按组派单；Owner/成员权限受控；跨产品组访问被拒。",
            0.75, "已完成",
            "服务层、企业端和组织页回归通过；未改写工作树中已有的组织页面修改。",
            "Go 产品组/权限/分派回归及 Node 组织页 10 项测试：PASS。",
            "正式服务目录、专家联系人、合同和 OLA 数据仍由业务方提供。",
        ],
        [
            "PLAN-005", date(2026, 9, 11), "", "汇总回归与交付记录",
            "运行相关后端、配置、企业接口、组织页和 TypeScript 检查，记录结果与剩余风险。",
            "定向回归和类型检查通过；未完成的宽回归如实记录。",
            0.25, "已完成",
            "相关定向测试全部通过；完整 services 包回归运行超过 4 分钟无输出后主动中止。",
            "见“验证记录”工作表。",
            "业务方 UAT、真实外部账号和生产安全扫描不在本次本地验证范围。",
        ],
    ]
    for row in tasks:
        ws.append(row)

    style_header(ws, 6, 1, 11)
    style_body(ws, 7, 11, 1, 11)
    for row in range(7, 12):
        ws.cell(row, 2).number_format = "yyyy-mm-dd"
        ws.cell(row, 7).number_format = '0.00 "人日"'
        ws.row_dimensions[row].height = 82
    ws.freeze_panes = "A7"
    add_table(ws, "A6:K11", "WeeklyTasks")
    set_widths(ws, {
        "A": 13, "B": 13, "C": 12, "D": 24, "E": 42, "F": 46,
        "G": 12, "H": 12, "I": 46, "J": 38, "K": 42,
    })
    status_validation = DataValidation(type="list", formula1='"待处理,进行中,已完成,受阻"', allow_blank=False)
    ws.add_data_validation(status_validation)
    status_validation.add("H7:H100")
    ws.conditional_formatting.add(
        "H7:H100",
        FormulaRule(formula=['H7="已完成"'], fill=PatternFill("solid", fgColor=GREEN), font=Font(color=GREEN_TEXT)),
    )
    ws.conditional_formatting.add(
        "H7:H100",
        FormulaRule(formula=['H7="受阻"'], fill=PatternFill("solid", fgColor=RED)),
    )

    basis = wb.create_sheet("排期依据")
    base_sheet(
        basis,
        "排期依据",
        "将原修订讨论稿中的事实、假设和本周取舍分开记录，避免把讨论日期当成交付承诺。",
        5,
    )
    basis.append([])
    basis.append(["项目", "原表信息", "本周判断", "来源/证据", "备注"])
    basis_rows = [
        ["功能总数", 90, "保持原范围，不尝试本周覆盖全部功能", "排期总览、功能完成排期", "原文覆盖和运行验收是两件事。"],
        ["实现判断", "部分完成 47；待完成 28；待复验 12；待验证 3", "优先完成 3 个待验证项", "功能完成排期 D 列", "本周选择 FUN-021、FUN-028、FUN-083。"],
        ["S0 原拟日期", "2026-09-14 至 2026-09-25", "本周作为 S0 前置验证周", "迭代计划 S0", "日期未获团队确认。"],
        ["S0 原拟范围", "9 个功能项，每项原估 3 人日", "只承诺可独立验证的 3 项", "迭代计划、功能完成排期", "其余项目受身份源、SLA、专家或 Bot 规则约束。"],
        ["M0 原拟目标", "2026-09-25：现状验证、范围和验收口径冻结", "本周先产生可复验技术证据", "排期总览 M0", "业务签收仍需负责人。"],
        ["总体风险", "889 人日和 21 个双周迭代均待重估", "不以本周结果推导整体完工日", "排期总览第 11、28、30 行", "需真实团队容量和依赖日期。"],
    ]
    for row in basis_rows:
        basis.append(row)
    style_header(basis, 4, 1, 5)
    style_body(basis, 5, 10, 1, 5)
    for row in range(5, 11):
        basis.row_dimensions[row].height = 56
    basis.freeze_panes = "A5"
    add_table(basis, "A4:E10", "ScheduleBasis")
    set_widths(basis, {"A": 22, "B": 42, "C": 42, "D": 34, "E": 40})

    checks = wb.create_sheet("验证记录")
    base_sheet(
        checks,
        "验证记录",
        "记录本次执行的可重复检查。PASS 表示本地自动化通过，不代表业务方 UAT 或生产环境验收。",
        6,
    )
    checks.append([])
    checks.append(["编号", "范围", "命令/检查", "结果", "用时/数量", "说明"])
    check_rows = [
        ["CHK-001", "FUN-021 + FUN-083 服务层", "go test ./internal/services -run 'Test(TicketAutoClose|InspectUpload)' -count=1", "PASS", "6.355s", "新增边界和安全用例通过。"],
        ["CHK-002", "FUN-028 服务层", "go test ./internal/services -run '<产品组/分派/权限相关用例>' -count=1", "PASS", "1.259s", "覆盖跨组隔离、多产品组、分派和成员权限。"],
        ["CHK-003", "FUN-028 企业接口 + FUN-083 上传集成", "go test ./internal/handlers/enterprise -run '<相关企业接口用例>' -count=1", "PASS", "1.001s", "覆盖 Owner 转移、产品范围和文档上传。"],
        ["CHK-004", "配置", "go test ./internal/pkg/config -count=1", "PASS", "0.612s", "新增归档安全配置未破坏配置包。"],
        ["CHK-005", "组织页", "node --test app/enterprise/org/member-name.test.mjs", "PASS", "10/10", "产品组成员选择、名称策略和组织视图通过。"],
        ["CHK-006", "前端类型", "tsc --noEmit", "PASS", "完成", "直接使用已安装的 TypeScript 编译器检查。"],
        ["CHK-007", "完整 Go 包宽回归", "go test ./internal/services ./internal/handlers/enterprise ./internal/pkg/config -count=1", "未完成", ">4 分钟", "持续无输出后主动中止；不影响上述定向回归结论。"],
        ["CHK-008", "代码差异", "git diff --check", "PASS", "完成", "未发现空白或补丁格式错误。"],
    ]
    for row in check_rows:
        checks.append(row)
    style_header(checks, 4, 1, 6)
    style_body(checks, 5, 12, 1, 6)
    for row in range(5, 13):
        checks.row_dimensions[row].height = 55
    checks.freeze_panes = "A5"
    add_table(checks, "A4:F12", "VerificationLog")
    set_widths(checks, {"A": 12, "B": 34, "C": 72, "D": 14, "E": 16, "F": 46})
    checks.conditional_formatting.add(
        "D5:D100",
        FormulaRule(formula=['D5="PASS"'], fill=PatternFill("solid", fgColor=GREEN), font=Font(color=GREEN_TEXT)),
    )
    checks.conditional_formatting.add(
        "D5:D100",
        FormulaRule(formula=['D5="未完成"'], fill=PatternFill("solid", fgColor=AMBER)),
    )

    pending = wb.create_sheet("待确认事项")
    base_sheet(
        pending,
        "待确认事项",
        "这些输入不阻塞本周代码交付，但会阻塞正式业务验收、生产启用或整体排期冻结。",
        6,
    )
    pending.append([])
    pending.append(["优先级", "关联功能", "待确认内容", "建议负责人", "需要时间", "影响"])
    pending_rows = [
        ["高", "FUN-021", "批准自动关闭天数、例外类型、关闭前提醒和重开服务目标", "服务运营负责人", "S0 验收前", "无法形成正式关闭策略签收。"],
        ["高", "FUN-083/FUN-084", "批准文件类型、大小、ZIP 限制；确认隔离区和病毒处置流程", "客户 IT/安全", "S0 验收前", "只能完成技术安全基线，不能关闭安全验收。"],
        ["高", "FUN-028", "提供正式产品 Owner、成员和产品组映射", "产品/服务运营负责人", "S0 验收前", "当前只能用测试数据验证配置能力。"],
        ["高", "S0 其余功能", "统一身份测试账号、SLA 规则、Bot 规则、专家及 OLA 数据", "客户 IT/业务负责人", "开发前", "受依赖项无法承诺本周完成。"],
        ["中", "整体计划", "确认真实团队、每人投入、职责、节假日和并行容量", "项目负责人", "计划冻结前", "无法确认 889 人日和里程碑日期。"],
        ["中", "总体架构", "确认遵循原文组件/事实源，或批准当前替代方案", "架构/项目负责人", "计划冻结前", "架构差异可能产生返工。"],
    ]
    for row in pending_rows:
        pending.append(row)
    style_header(pending, 4, 1, 6)
    style_body(pending, 5, 10, 1, 6)
    for row in range(5, 11):
        pending.row_dimensions[row].height = 62
    pending.freeze_panes = "A5"
    add_table(pending, "A4:F10", "PendingItems")
    set_widths(pending, {"A": 12, "B": 22, "C": 56, "D": 28, "E": 18, "F": 44})

    for sheet in wb.worksheets:
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_view.zoomScale = 90
        sheet.auto_filter.ref = sheet.tables[next(iter(sheet.tables))].ref if sheet.tables else None

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_workbook()
