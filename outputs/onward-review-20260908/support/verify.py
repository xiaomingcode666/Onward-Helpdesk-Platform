import collections
import json
from pathlib import Path

import openpyxl

folder = Path('D:/work/helpdesk/outputs/onward-review-20260908')
source = Path('D:/桌面/Onward_业务功能对照与实施排期_最终版_20260907.xlsx')
output = folder / 'Onward_业务功能对照与实施排期_修订讨论稿_20260908.xlsx'
baseline = openpyxl.load_workbook(source, data_only=False)
wb = openpyxl.load_workbook(output, data_only=False)
cached = openpyxl.load_workbook(output, data_only=True)
spec = json.loads((folder / 'support/verification.json').read_text(encoding='utf8'))
assert wb.sheetnames[:5] == baseline.sheetnames
assert len(wb.sheetnames) == 7
main = wb['功能完成排期']
assert main.max_row == 91
assert len({main.cell(r, 1).value for r in range(2, 92)}) == 90
for row in range(2, 92):
    feature_id = main.cell(row, 1).value
    assert feature_id == baseline['功能完成排期'].cell(row, 1).value
    for column in [10, 11, 12, 17, 18]:
        before = baseline['功能完成排期'].cell(row, column).value
        after = main.cell(row, column).value
        assert before == after, (feature_id, column, before, after)
    assert main.cell(row, 4).value == spec['expected'][feature_id]['status']
    assert main.cell(row, 16).value == '；'.join(spec['expected'][feature_id]['refs'])
    assert main.cell(row, 23).value == '未确认'
    assert main.cell(row, 22).value == '未完成端到端验收'
    assert 'DES-' not in main.cell(row, 16).value
counts = collections.Counter(main.cell(r, 4).value for r in range(2, 92))
assert counts == {k: v for k, v in spec['counts'].items() if v}
assert cached['排期总览']['A5'].value == 90
assert cached['排期总览']['C5'].value == 0
assert cached['排期总览']['G7'].value == 889
assert sum(cached['迭代计划'].cell(r, 7).value for r in range(2, 23)) == 889
assert sum(cached['迭代计划'].cell(r, 6).value for r in range(2, 23)) == 78
for r in range(2, 80):
    feature_id = cached['优先执行清单'].cell(r, 2).value
    source_row = spec['expected'][feature_id]['row']
    for a, b in [(3,2),(4,3),(5,4),(6,7),(7,9),(8,11),(9,6),(10,15)]:
        assert cached['优先执行清单'].cell(r,a).value == cached['功能完成排期'].cell(source_row,b).value, (feature_id,a,b)
for name in ['迭代计划','外部准备项']:
    for r in range(2, baseline[name].max_row+1):
        columns = [3,4] if name == '迭代计划' else [3,7]
        for c in columns:
            assert wb[name].cell(r,c).value == baseline[name].cell(r,c).value, (name,r,c)
source_ids = {wb['需求原文与范围'].cell(r,1).value for r in range(5,152)}
for values in spec['expected'].values():
    assert set(values['refs']) <= source_ids
assert 'FND-001' not in main['P80'].value
assert wb['逐项核查依据'].max_row == 94
assert {wb['逐项核查依据'].cell(r,1).value for r in range(5,95)} == set(spec['expected'])
errors = []
for sheet in wb:
    for row in sheet:
        for cell in row:
            if cell.data_type == 'e' or cell.value == '[object Object]':
                errors.append((sheet.title,cell.coordinate,cell.value))
    for table in sheet.tables.values():
        first_row = openpyxl.utils.range_boundaries(table.ref)[1]
        assert table.autoFilter and table.autoFilter.ref == table.ref
        for i,column in enumerate(table.tableColumns,1):
            assert column.name == sheet.cell(first_row,i).value, (sheet.title,column.name,sheet.cell(first_row,i).value)
    assert max((d.height or 0) for d in sheet.row_dimensions.values()) <= 409
assert not errors, errors
print(json.dumps({'result':'PASS','sheets':len(wb.sheetnames),'features':90,'source_entries':147,'original_dates_effort_owners_preserved':True,'priority_and_iteration_reconciled':True,'tables_with_filters':sum(len(s.tables) for s in wb),'status_counts':dict(counts),'size_bytes':output.stat().st_size},ensure_ascii=False))
